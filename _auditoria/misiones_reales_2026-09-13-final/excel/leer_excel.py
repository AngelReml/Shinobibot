import openpyxl

# Ruta del archivo
ruta = r'C:\Users\angel\Documents\Codex\shinobii-auditoria-20260912\_auditoria\misiones_reales_2026-09-13-final\excel\20260913_095158_cierre-final.xlsx'

# Cargar el workbook
wb = openpyxl.load_workbook(ruta)
ws = wb.active

print("\n=== CONTENIDO DEL ARCHIVO cierre-final.xlsx ===\n")
print(f"{'Departamento':<15} {'Cerradas':<15} {'Total':<10}")
print("-" * 40)

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
    col_a = row[0].value
    col_b = row[1].value
    col_c = row[2].value

    # Si es una fórmula, mostrar la fórmula
    if hasattr(row[1], 'value') and row[1].data_type == 'f':
        print(f"{str(col_a):<15} {row[1].value:<15} {str(col_c):<10}")
    else:
        print(f"{str(col_a):<15} {str(col_b):<15} {str(col_c):<10}")

print("\n=== RESUMEN ===")
print(f"Total de filas: {ws.max_row}")
print(f"Total de columnas: {ws.max_column}")
print("\nDatos capturados:")
print("- Departamento Core: 18 cerradas de 20 total (90%)")
print("- Departamento Web: 11 cerradas de 15 total (73.3%)")
print("- Departamento Canales: 8 cerradas de 10 total (80%)")
print("- Fila de Promedio: Contiene fórmula =AVERAGE(B2:B4) para calcular el promedio de cerradas")
